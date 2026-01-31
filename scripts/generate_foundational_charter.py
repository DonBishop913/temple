"""Generate a Word .docx foundational charter from the markdown template.

Usage:
    python scripts/generate_foundational_charter.py --name "Donald M. Miller" --address "[Your Address, Overland Park, KS 66212]" --date 2025-10-20

This script requires `python-docx` (pip install python-docx).
"""
import argparse
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except Exception:
    Document = None


TEMPLATE_MD = Path('foundational_charter/TempleCharter_CouncilSeal_Packet.md')
OUT_DIR = Path('docs')
OUT_DOCX = OUT_DIR / 'foundational_charter.docx'


def simple_md_to_doc(document, md_text: str):
    # Very small markdown-ish renderer: headings and paragraphs
    for line in md_text.splitlines():
        line = line.rstrip()
        if not line:
            document.add_paragraph('')
            continue
        if line.startswith('# '):
            p = document.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith('## '):
            document.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith('### '):
            document.add_heading(line[4:].strip(), level=3)
            continue
        # simple list handling
        if line.lstrip().startswith('- '):
            document.add_paragraph(line.lstrip()[2:].strip(), style='List Bullet')
            continue
        # normal paragraph
        p = document.add_paragraph(line)
        # small font tweak
        for run in p.runs:
            run.font.size = Pt(11)


def generate(name: str, address: str, date: str):
    if Document is None:
        raise SystemExit('python-docx is not installed. Run: pip install python-docx')

    if not TEMPLATE_MD.exists():
        raise SystemExit(f'Template not found: {TEMPLATE_MD}')

    md = TEMPLATE_MD.read_text(encoding='utf-8')
    # Replace obvious placeholders
    md = md.replace('[Your Address, Overland Park, KS 66212]', address)
    md = md.replace('[YYYY-MM-DD]', date)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    # Some static analyzers can't see python-docx attributes; guard access and keep runtime behavior.
    try:
        doc.styles['Normal'].font.name = 'Calibri'  # type: ignore[attr-defined]
        doc.styles['Normal'].font.size = Pt(11)     # type: ignore[attr-defined]
    except Exception:
        # If the runtime python-docx version doesn't expose these attributes, continue without failing.
        pass

    # Title
    doc.add_heading('Articles of Incorporation for Unified Sovereign Intelligence Coalition (USIC)', level=1)

    # Add the rendered template
    simple_md_to_doc(doc, md)

    # Signature block
    doc.add_paragraph('')
    doc.add_paragraph(f'Signed: {name}')
    doc.add_paragraph(f'Date: {date}')

    # Ensure path is a str for python-docx save compatibility
    doc.save(str(OUT_DOCX))
    print('Wrote', OUT_DOCX)


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--name', default='Donald M. Miller')
    p.add_argument('--address', default='[Your Address, Overland Park, KS 66212]')
    p.add_argument('--date', default='2025-10-20')
    args = p.parse_args()
    generate(args.name, args.address, args.date)


if __name__ == '__main__':
    main()
